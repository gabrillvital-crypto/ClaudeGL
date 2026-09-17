"""
Plano de Ação BPO — Vinci Airports
Versão 3 — 16/09/2026
5 slides: Capa + Contexto + Etapas 1-3 + Etapas 4-5 + Cronograma/Próximos passos
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─────────────────────────────────────────────────────────────
# Paleta Efcaz
# ─────────────────────────────────────────────────────────────
TEAL       = RGBColor(0x0E, 0x8F, 0xA3)
TEAL_DARK  = RGBColor(0x09, 0x6A, 0x7A)
TEAL_PALE  = RGBColor(0xE6, 0xF5, 0xF7)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY       = RGBColor(0x55, 0x65, 0x70)
DARK       = RGBColor(0x17, 0x20, 0x2A)
OFFWHITE   = RGBColor(0xF6, 0xFB, 0xFC)
BORDER_CLR = RGBColor(0xB2, 0xDE, 0xE5)

# ─────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────

def rect(slide, x, y, w, h, fill_rgb, border_rgb=None, border_pt=0.75):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill_rgb
    if border_rgb:
        s.line.color.rgb = border_rgb
        s.line.width = Pt(border_pt)
    else:
        s.line.fill.background()
    return s


def tb(slide, text, x, y, w, h, size=11, bold=False,
       color=DARK, align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def tb_lines(slide, lines, x, y, w, h, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get('align', PP_ALIGN.LEFT)
        sb = ln.get('space_before', 0)
        if sb:
            p.space_before = Pt(sb)
        run = p.add_run()
        run.text = ln['text']
        run.font.size = Pt(ln.get('size', 11))
        run.font.bold = ln.get('bold', False)
        run.font.color.rgb = ln.get('color', DARK)
    return box


# ─────────────────────────────────────────────────────────────
# Componentes reutilizáveis
# ─────────────────────────────────────────────────────────────

def header_band(slide, title, subtitle):
    rect(slide, 0, 0, 10, 0.95, TEAL)
    tb(slide, title,    0.35, 0.07, 9.3, 0.47, size=21, bold=True,  color=WHITE)
    tb(slide, subtitle, 0.35, 0.57, 9.3, 0.34, size=9.5, bold=False, color=WHITE)


def metric_card(slide, x, y, w, h, label, value):
    rect(slide, x, y, w, h, TEAL_PALE, TEAL, border_pt=1.0)
    tb(slide, label, x+0.16, y+0.18, w-0.28, 0.42,
       size=8.5, bold=False, color=GRAY)
    tb(slide, value, x+0.16, y+0.63, w-0.28, 1.2,
       size=16, bold=True, color=TEAL)


def action_card(slide, x, y, w, h, number, title, bullets):
    rect(slide, x, y, w, h, OFFWHITE, BORDER_CLR, border_pt=0.75)
    rect(slide, x, y, 0.07, h, TEAL)

    bs = 0.40
    bx, by = x + 0.16, y + 0.18
    rect(slide, bx, by, bs, bs, TEAL)
    tb(slide, number, bx, by + 0.03, bs, bs - 0.04,
       size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    tb(slide, title, x + 0.65, y + 0.20, w - 0.75, 0.55,
       size=10.5, bold=True, color=TEAL_DARK)

    rect(slide, x + 0.16, y + 0.77, w - 0.26, 0.01, BORDER_CLR)

    if bullets:
        lines = []
        for i, b in enumerate(bullets):
            lines.append({
                'text': '  ' + b,
                'size': 9,
                'bold': False,
                'color': GRAY,
                'align': PP_ALIGN.LEFT,
                'space_before': 5 if i > 0 else 0,
            })
        tb_lines(slide, lines, x + 0.18, y + 0.83, w - 0.28, h - 0.93)


# ─────────────────────────────────────────────────────────────
# Criar apresentação
# ─────────────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width  = Inches(10)
prs.slide_height = Inches(5.62)

blank = next(l for l in prs.slide_layouts if l.name == 'Blank')


# ═════════════════════════════════════════════════════════════
# SLIDE 1 — CAPA
# ═════════════════════════════════════════════════════════════

sl1 = prs.slides.add_slide(blank)
sl1.background.fill.solid()
sl1.background.fill.fore_color.rgb = TEAL_DARK

# Faixa decorativa inferior
rect(sl1, 0, 4.45, 10, 1.17, TEAL)

# Título
tb(sl1, 'Saneamento da Base\nde Fornecedores',
   0.85, 1.15, 8.3, 2.0, size=33, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

# Linha divisória
rect(sl1, 0.85, 3.28, 3.8, 0.04, TEAL_PALE)

# Subtítulo
tb(sl1, 'Plano de Ação — BPO Documental',
   0.85, 3.42, 8.0, 0.5, size=13, bold=False, color=WHITE, align=PP_ALIGN.LEFT)

# Cliente e data
tb(sl1, 'Vinci Airports  \u2022  Setembro / 2026',
   0.85, 3.94, 8.0, 0.4, size=10.5, bold=False, color=TEAL_PALE, align=PP_ALIGN.LEFT)


# ═════════════════════════════════════════════════════════════
# SLIDE 2 — CONTEXTO DO PROJETO
# ═════════════════════════════════════════════════════════════

sl2 = prs.slides.add_slide(blank)
sl2.background.fill.solid()
sl2.background.fill.fore_color.rgb = WHITE

header_band(sl2,
    'Contexto do Projeto',
    'Levantamento de necessidades  \u2022  Vinci Airports  \u2022  Set/2026')

cx   = [0.28, 3.44, 6.60]
cy   = [1.08, 3.18]
cw, ch = 3.0, 1.97

metrics = [
    ('Fornecedores na base (SAP)',        '5.634'),
    ('Contrato Efcaz',                    '500 fornecedores'),
    ('Fornecedores ativos em operacao',   '300 ativos'),
    ('Perguntas - Reforma Tributaria',    '14 perguntas'),
    ('Campos novos nao mapeados',         '~6 campos'),
    ('Prazo de saneamento',               'Antes da calculadora KPMG'),
]

for i, (label, value) in enumerate(metrics):
    row, col = divmod(i, 3)
    metric_card(sl2, cx[col], cy[row], cw, ch, label, value)


# ═════════════════════════════════════════════════════════════
# SLIDE 3 — PLANO DE AÇÃO: ETAPAS 1, 2 e 3
# ═════════════════════════════════════════════════════════════

sl3 = prs.slides.add_slide(blank)
sl3.background.fill.solid()
sl3.background.fill.fore_color.rgb = WHITE

header_band(sl3,
    'Plano de Ação — Etapas 1, 2 e 3',
    'Saneamento da Base  \u2022  Vinci Airports  \u2022  Set/2026')

ax = [0.28, 3.44, 6.60]
ay, aw, ah = 1.07, 3.0, 4.37

action_card(sl3, ax[0], ay, aw, ah,
    '01',
    'Estruturacao na plataforma',
    [
        'Criacao de linha de fornecimento temporaria, sem documentos e sem buscas automaticas.',
        'Consulta de CNPJ (Receita Federal) para triagem de status ativo/inativo.',
        'Buscas automaticas completas ativadas apos a conclusao do saneamento.',
    ])

action_card(sl3, ax[1], ay, aw, ah,
    '02',
    'Campos dinamicos obrigatorios',
    [
        '~6 novos campos da Reforma Tributaria incluidos como campos dinamicos.',
        'Dados bancarios (conta e PIX) obrigatorios na ficha de cadastro.',
        'Fornecedor nao consegue submeter sem preencher todos os campos.',
    ])

action_card(sl3, ax[2], ay, aw, ah,
    '03',
    'Convite BPO - novo cadastro',
    [
        'Fornecedores convidados a se cadastrar como novos — nao apenas a atualizar.',
        'Rastreamento via RD Station: envio, abertura e cadastro monitorados.',
        'Ate 3 tentativas por CNPJ, com lembretes semanais.',
        'CNPJs sem retorno reportados a Vinci para decisao sobre tratativa.',
    ])


# ═════════════════════════════════════════════════════════════
# SLIDE 4 — PLANO DE AÇÃO: ETAPAS 4 e 5
# ═════════════════════════════════════════════════════════════

sl4 = prs.slides.add_slide(blank)
sl4.background.fill.solid()
sl4.background.fill.fore_color.rgb = WHITE

header_band(sl4,
    'Plano de Ação — Etapas 4 e 5',
    'Saneamento da Base  \u2022  Vinci Airports  \u2022  Set/2026')

action_card(sl4, ax[0], ay, aw, ah,
    '04',
    'Validacao bancaria',
    [
        'Validacao de conta bancaria e chave PIX como criterio de bloqueio para submissao.',
        'Fornecedor sem dado bancario validado nao avanca no fluxo de cadastro.',
        'Validacao realizada conforme os fornecedores preenchem os dados na plataforma.',
    ])

action_card(sl4, ax[1], ay, aw, ah,
    '05',
    'Analise e aprovacao',
    [
        'Equipe analisa cada submissao individualmente.',
        'Aprovado: vinculado a linha definitiva e homologacao completa.',
        'Aprovado com ressalva: pendencia gerada para complemento.',
        'Reprovado definitivamente: fornecedor inativado (historico preservado).',
    ])

action_card(sl4, ax[2], ay, aw, ah,
    '+',
    'Criterios de aprovacao',
    [
        'Definidos em conjunto com a Vinci — nao de forma unilateral.',
        'Referencia: parametros em alinhamento com a KPMG.',
        'Consistencia entre o processo de saneamento e as exigencias da Reforma Tributaria.',
    ])


# ═════════════════════════════════════════════════════════════
# SLIDE 5 — CRONOGRAMA E PRÓXIMOS PASSOS
# ═════════════════════════════════════════════════════════════

sl5 = prs.slides.add_slide(blank)
sl5.background.fill.solid()
sl5.background.fill.fore_color.rgb = WHITE

header_band(sl5,
    'Cronograma e Próximos Passos',
    'Saneamento da Base  \u2022  Vinci Airports  \u2022  Set/2026')

action_card(sl5, 0.28, 1.07, 4.5, 4.37,
    'T',
    'Cronograma',
    [
        '8 semanas a partir do OK formal e entrega da base de contatos.',
        'Inicio previsto: 2a quinzena de setembro/2026.',
        'Entrega estimada: novembro/2026.',
        'Reunioes quinzenais de acompanhamento com Vinci.',
        'Relatorios de evolucao disponiveis a qualquer momento: envios x cadastros x aprovacoes.',
    ])

action_card(sl5, 5.22, 1.07, 4.5, 4.37,
    'P',
    'Para iniciarmos',
    [
        'Criterios de aprovacao e reprovacao (em alinhamento com a KPMG).',
        'Planilha de fornecedores com e-mails de contato atualizados.',
        'OK formal da Vinci para inicio do projeto.',
    ])


# ─────────────────────────────────────────────────────────────
# Salvar
# ─────────────────────────────────────────────────────────────

output = (
    r"C:\Users\gabriel.evangelista\Documents\ClaudeGL\clientes\vinci_airports"
    r"\VinciAirports_PlanoAcao_BPO_16-09-2026.pptx"
)
prs.save(output)
print(f"Arquivo salvo: {output}")
print(f"Slides gerados: {len(prs.slides)}")
