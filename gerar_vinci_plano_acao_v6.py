"""
gerar_vinci_plano_acao_v6.py
Melhorias textuais no VinciAirports_PlanoAcao_BPO_16-09-2026_v2.pptx
Saída: VinciAirports_PlanoAcao_BPO_16-09-2026_v3.pptx
"""

import sys
import io
import copy
from pptx import Presentation

# Força UTF-8 no stdout para evitar UnicodeEncodeError no Windows
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pptx.util import Pt
from pptx.dml.color import RGBColor

INPUT  = r"clientes\vinci_airports\VinciAirports_PlanoAcao_BPO_16-09-2026_v2.pptx"
OUTPUT = r"clientes\vinci_airports\VinciAirports_PlanoAcao_BPO_16-09-2026_v3.pptx"

# ---------------------------------------------------------------------------
# Mapeamento de substituições simples (texto exato → novo texto)
# Chave = texto COMPLETO do parágrafo (para localizar o shape certo)
# Valor = novo texto (aplicado ao primeiro run do parágrafo)
# ---------------------------------------------------------------------------
SUBSTITUTIONS = {
    # Slide 1 — Capa
    "Vinci Airports  •  Setembro / 2026": "Vinci Airports  •  Setembro/2026",

    # Slide 2 — Contexto: KPI "300 ativos" repete palavra da label
    "300 ativos": "300",

    # Slide 2 — Label mais clara para campo novo
    "Campos novos não mapeados": "Novos campos identificados",

    # Slide 3 — Título da Etapa 03 mais preciso
    "Convite BPO — novo cadastro": "Convite BPO: recadastro completo",

    # Slide 6 — Mensagem de encerramento mais direta e orientada à ação
    "Estamos prontos para iniciar assim que recebermos o alinhamento da Vinci.": (
        "Aguardamos o OK da Vinci para iniciar o saneamento da base."
    ),
}

# Slide 5 — bullet com × como separador: substituição parcial dentro do run
PARTIAL_REPLACE = {
    "envios × cadastros × aprovações": "envios, cadastros e aprovações",
}

# ---------------------------------------------------------------------------

def apply_substitutions(prs: Presentation) -> list[str]:
    """Percorre todos os shapes e aplica as substituições. Retorna log."""
    log = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full_text = para.text

                # --- Substituição exata de parágrafo inteiro ---
                if full_text in SUBSTITUTIONS:
                    new_text = SUBSTITUTIONS[full_text]
                    # Preserva formatação: aplica no 1º run, limpa os demais
                    if para.runs:
                        para.runs[0].text = new_text
                        for run in para.runs[1:]:
                            run.text = ""
                    log.append(
                        f"  [SLIDE {slide_idx}] '{shape.name}' "
                        f"'{full_text[:60]}' → '{new_text[:60]}'"
                    )

                # --- Substituição parcial dentro de um run ---
                else:
                    for run in para.runs:
                        for old, new in PARTIAL_REPLACE.items():
                            if old in run.text:
                                run.text = run.text.replace(old, new)
                                log.append(
                                    f"  [SLIDE {slide_idx}] '{shape.name}' "
                                    f"parcial: '{old}' → '{new}'"
                                )
    return log


def verify_no_overflow(prs: Presentation):
    """Aviso se algum parágrafo tiver mais de 100 caracteres (possível overflow)."""
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = para.text
                if len(t) > 100:
                    print(f"  [AVISO OVERFLOW] Slide {slide_idx} / {shape.name}: {len(t)} chars — '{t[:80]}...'")


def main():
    print(f"Abrindo: {INPUT}")
    prs = Presentation(INPUT)

    print("\nAplicando melhorias textuais...")
    log = apply_substitutions(prs)
    if log:
        for entry in log:
            print(entry)
    else:
        print("  Nenhuma substituição encontrada — verifique o mapeamento.")

    print("\nVerificando overflow de texto...")
    verify_no_overflow(prs)

    print(f"\nSalvando em: {OUTPUT}")
    prs.save(OUTPUT)
    print("✓ Arquivo salvo com sucesso.")

    # Confirmação rápida lendo de volta
    print("\n--- Confirmação (textos alterados) ---")
    prs2 = Presentation(OUTPUT)
    targets = set(SUBSTITUTIONS.values()) | set(PARTIAL_REPLACE.values())
    for slide_idx, slide in enumerate(prs2.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                for target in targets:
                    if target in t:
                        print(f"  [OK] Slide {slide_idx} / {shape.name}: '{t[:80]}'")


if __name__ == "__main__":
    main()
